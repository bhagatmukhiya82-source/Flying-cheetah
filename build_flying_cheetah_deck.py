import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_PPTX = r"c:\AUTOMATION BUSSNIESS\Flying_Cheetah_Autonomous_Commerce_Pitch_Deck.pptx"
ASSETS_DIR = r"c:\AUTOMATION BUSSNIESS\docs\assets"
IMAGES_DIR = r"c:\AUTOMATION BUSSNIESS\flying-cheetah\images"

# -------------------------------------------------------------
# PALETTE DEFINITIONS (NANO BANANA CYBERPUNK)
# -------------------------------------------------------------
C_BG = RGBColor(5, 7, 16)             # Deep Obsidian #050710
C_CARD_BG = RGBColor(15, 20, 36)      # Dark Tech Slate Navy
C_CARD_BORDER = RGBColor(38, 52, 82)  # Subtle Slate Border
C_GOLD = RGBColor(255, 230, 0)        # #FFE600
C_AMBER = RGBColor(255, 160, 0)       # #FFA000
C_CYAN = RGBColor(0, 240, 255)        # #00F0FF
C_MAGENTA = RGBColor(255, 0, 127)     # #FF007F
C_EMERALD = RGBColor(0, 255, 136)     # #00FF88
C_VIOLET = RGBColor(168, 85, 247)     # #A855F7
C_WHITE = RGBColor(255, 255, 255)
C_SLATE = RGBColor(160, 175, 200)     # Body Text
C_MUTED = RGBColor(110, 125, 150)
C_RED = RGBColor(244, 63, 94)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"

def set_slide_background(slide, prs):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_BG
    bg.line.fill.background()
    return bg

def add_header(slide, tag_text, title_text, sub_text, tag_color=C_GOLD):
    # Category Tag / Breadcrumb
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.35))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.font.name = FONT_HEADING
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = tag_color

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.5), Inches(0.65))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(23)
    p_t.font.bold = True
    p_t.font.color.rgb = C_WHITE

    # Subtitle
    if sub_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.5), Inches(0.4))
        tf_s = sub_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        p_s.text = sub_text
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(12.5)
        p_s.font.color.rgb = C_SLATE

def add_card(slide, left, top, width, height, border_color=C_CARD_BORDER, fill_color=C_CARD_BG):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = fill_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.5)
    return card

def safe_add_image(slide, path, left, top, width=None, height=None):
    if os.path.exists(path):
        try:
            if width and height:
                return slide.shapes.add_picture(path, left, top, width=width, height=height)
            elif width:
                return slide.shapes.add_picture(path, left, top, width=width)
            elif height:
                return slide.shapes.add_picture(path, left, top, height=height)
            else:
                return slide.shapes.add_picture(path, left, top)
        except Exception as e:
            print(f"Warning: could not add image {path}: {e}")
    else:
        print(f"Warning: Image file not found: {path}")
    return None

def build_presentation():
    prs = pptx.Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    print("--- BUILDING 13-SLIDE MASTER PRESENTATION ---")

    # =========================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, prs)

    add_card(s1, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9), border_color=C_GOLD)

    # Logo
    logo_path = os.path.join(IMAGES_DIR, "flying_cheetah_hd_logo_crest.jpg")
    if not os.path.exists(logo_path):
        logo_path = os.path.join(ASSETS_DIR, "logo_heraldic.jpg")
    safe_add_image(s1, logo_path, Inches(1.3), Inches(1.3), height=Inches(1.5))

    # Title & Tagline Box
    tb1 = s1.shapes.add_textbox(Inches(3.1), Inches(1.2), Inches(8.8), Inches(2.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "FLYING CHEETAH LABS"
    p.font.name = FONT_HEADING
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = C_GOLD

    p2 = tf1.add_paragraph()
    p2.text = "All-in-One AI Marketing & Autonomous Commerce Platform"
    p2.font.name = FONT_HEADING
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = C_WHITE

    p3 = tf1.add_paragraph()
    p3.text = "Empowering retail, D2C, and local businesses with human-like targeted marketing at scale and zero-payroll hyperlocal fleet dispatch."
    p3.font.name = FONT_BODY
    p3.font.size = Pt(13)
    p3.font.color.rgb = C_SLATE

    # 3 Stat Pills
    metrics = [
        ("3-TAP COMMERCE", "WhatsApp native catalog checkout", C_CYAN),
        ("SUB-600MS AI VOICE", "Zero-payroll COD & support bots", C_EMERALD),
        ("1-CLICK DISPATCH", "Borzo & Porter bike integration", C_MAGENTA)
    ]
    for i, (m_title, m_desc, m_color) in enumerate(metrics):
        mx = Inches(1.3 + i * 3.75)
        add_card(s1, mx, Inches(3.6), Inches(3.5), Inches(1.4), border_color=m_color)
        tb_m = s1.shapes.add_textbox(mx + Inches(0.2), Inches(3.75), Inches(3.1), Inches(1.1))
        tf_m = tb_m.text_frame
        tf_m.word_wrap = True
        pm1 = tf_m.paragraphs[0]
        pm1.text = m_title
        pm1.font.name = FONT_HEADING
        pm1.font.size = Pt(14)
        pm1.font.bold = True
        pm1.font.color.rgb = m_color
        pm2 = tf_m.add_paragraph()
        pm2.text = m_desc
        pm2.font.name = FONT_BODY
        pm2.font.size = Pt(11)
        pm2.font.color.rgb = C_WHITE

    # Footer Leadership Bar
    add_card(s1, Inches(1.3), Inches(5.3), Inches(10.733), Inches(0.9), border_color=C_CARD_BORDER)
    tb_foot = s1.shapes.add_textbox(Inches(1.5), Inches(5.42), Inches(10.3), Inches(0.7))
    tf_foot = tb_foot.text_frame
    pf = tf_foot.paragraphs[0]
    pf.text = "Co-Founders: Bhagat Mukhiya & Yash Gaikwad   |   HQ: Opera House & Bandra West, Mumbai   |   Deployments: 2026-2027"
    pf.font.name = FONT_BODY
    pf.font.size = Pt(12)
    pf.font.bold = True
    pf.font.color.rgb = C_GOLD

    # =========================================================
    # SLIDE 2: THE MARKET BOTTLENECK
    # =========================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, prs)
    add_header(s2, "Market Reality & Operational Friction", "The 3 Fatal Bottlenecks Draining Local Retail & D2C Margins", 
               "Why traditional stores struggle to grow and why direct customer ownership is critical.", tag_color=C_RED)

    problems = [
        ("THE AGGREGATOR TRAP", "25% – 30% Margin Loss", C_RED, [
            "• Merchants surrender 25% to 30% of gross revenue to Swiggy, Zomato, and 3rd-party marketplaces.",
            "• Aggregators monopolize customer phone numbers and data, preventing direct retargeting.",
            "• Zero brand loyalty: customers belong to the platform, not to the local merchant."
        ]),
        ("THE AD SPEND BLACK HOLE", "40% Ad Spend Wasted", C_AMBER, [
            "• Leads from Instagram & Facebook go cold when response times exceed 5 minutes.",
            "• Manual staff cannot handle sudden inquiry spikes or midnight purchase intent.",
            "• WhatsApp messages left unreplied result in immediate competitor purchases."
        ]),
        ("THE DISPATCH NIGHTMARE", "30+ Mins Manual Hassle", C_MAGENTA, [
            "• Store managers spend 30+ minutes per order opening courier apps, entering addresses, and haggling.",
            "• Full-time delivery drivers require ₹18,000–₹22,000 monthly payrolls even on slow days.",
            "• Unconfirmed COD orders suffer 35%+ Return-to-Origin (RTO), destroying logistics margins."
        ])
    ]

    for i, (p_title, p_stat, p_col, p_bullets) in enumerate(problems):
        px = Inches(0.8 + i * 3.95)
        add_card(s2, px, Inches(1.9), Inches(3.8), Inches(5.0), border_color=p_col)
        
        # Pill Stat
        add_card(s2, px + Inches(0.3), Inches(2.2), Inches(3.2), Inches(0.7), border_color=p_col, fill_color=RGBColor(24, 20, 32))
        tb_ps = s2.shapes.add_textbox(px + Inches(0.35), Inches(2.28), Inches(3.1), Inches(0.5))
        tf_ps = tb_ps.text_frame
        p_stat_para = tf_ps.paragraphs[0]
        p_stat_para.text = p_stat
        p_stat_para.alignment = PP_ALIGN.CENTER
        p_stat_para.font.name = FONT_HEADING
        p_stat_para.font.size = Pt(14)
        p_stat_para.font.bold = True
        p_stat_para.font.color.rgb = p_col

        # Title
        tb_pt = s2.shapes.add_textbox(px + Inches(0.3), Inches(3.05), Inches(3.2), Inches(0.6))
        tf_pt = tb_pt.text_frame
        tf_pt.word_wrap = True
        p_t_para = tf_pt.paragraphs[0]
        p_t_para.text = p_title
        p_t_para.font.name = FONT_HEADING
        p_t_para.font.size = Pt(15)
        p_t_para.font.bold = True
        p_t_para.font.color.rgb = C_WHITE

        # Bullets
        tb_pb = s2.shapes.add_textbox(px + Inches(0.3), Inches(3.7), Inches(3.2), Inches(3.0))
        tf_pb = tb_pb.text_frame
        tf_pb.word_wrap = True
        for b in p_bullets:
            p_b = tf_pb.add_paragraph() if tf_pb.paragraphs[0].text else tf_pb.paragraphs[0]
            p_b.text = b
            p_b.font.name = FONT_BODY
            p_b.font.size = Pt(11)
            p_b.font.color.rgb = C_SLATE

    # =========================================================
    # SLIDE 3: THE AUTONOMOUS SOLUTION FLYWHEEL
    # =========================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, prs)
    add_header(s3, "Unified Architecture", "The Flying Cheetah Autonomous Commerce Flywheel", 
               "An interconnected 5-pillar ecosystem that eliminates friction from ad click to doorstep delivery.")

    # Left: Explanation Cards
    flywheel_steps = [
        ("1. Headless High-Speed Storefront", "Sub-second Next.js web catalog synced in real time.", C_CYAN),
        ("2. WhatsApp Conversational Commerce", "In-chat catalog, cart recovery, and 1-click UPI checkout.", C_EMERALD),
        ("3. Sub-600ms AI Voice Calling", "Autonomous voice agent confirms COD orders & verifies address.", C_GOLD),
        ("4. 1-Click Zero-Payroll Fleet Dispatch", "Instant Borzo bike assigned to store counter with 5-min prep buffer.", C_MAGENTA),
        ("5. Back-Office Gemini Vision OCR", "Automated Tally XML and Google Sheets sync for zero data entry.", C_VIOLET)
    ]

    for j, (step_t, step_d, step_c) in enumerate(flywheel_steps):
        sy = Inches(1.9 + j * 0.98)
        add_card(s3, Inches(0.8), sy, Inches(5.6), Inches(0.88), border_color=step_c)
        tb_s = s3.shapes.add_textbox(Inches(1.0), sy + Inches(0.1), Inches(5.2), Inches(0.68))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        p1 = tf_s.paragraphs[0]
        p1.text = step_t
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = step_c
        p2 = tf_s.add_paragraph()
        p2.text = step_d
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_WHITE

    # Right: Flywheel Image
    flywheel_img = os.path.join(ASSETS_DIR, "nano_banana_ecosystem_flywheel.png")
    if not os.path.exists(flywheel_img):
        flywheel_img = os.path.join(ASSETS_DIR, "ai_commerce_neural_matrix.png")
    add_card(s3, Inches(6.7), Inches(1.9), Inches(5.8), Inches(5.0), border_color=C_GOLD)
    safe_add_image(s3, flywheel_img, Inches(6.8), Inches(2.0), width=Inches(5.6), height=Inches(4.8))

    # =========================================================
    # SLIDE 4: PILLAR 1 - HIGH-SPEED STOREFRONTS
    # =========================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, prs)
    add_header(s4, "Pillar 1: Digital Front-Door", "Headless Storefronts Engineered for Ultra-Fast Conversions", 
               "Replacing bloated Shopify themes with blazing-fast Next.js architecture.", tag_color=C_CYAN)

    add_card(s4, Inches(0.8), Inches(1.9), Inches(5.8), Inches(5.0), border_color=C_CYAN)
    tb_p1 = s4.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.2), Inches(4.5))
    tf_p1 = tb_p1.text_frame
    tf_p1.word_wrap = True

    p = tf_p1.paragraphs[0]
    p.text = "SPEED EQUALS DIRECT REVENUE"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_CYAN

    bullets_p1 = [
        ("Sub-350ms Page Loads:", "Eliminates bounce rates caused by slow 5-second traditional e-commerce templates."),
        ("Direct WhatsApp Purchase Link:", "Every product card features an instant 'Order on WhatsApp' button with pre-filled SKU."),
        ("Instant Catalog Management:", "Merchants update prices, stock, and banners directly from a simplified mobile dashboard."),
        ("Full SEO & Schema Integration:", "Structured data for Google Local Search, rich snippets, and social previews."),
        ("Zero Server Maintenance:", "Built on high-availability serverless edge compute with 99.99% uptime.")
    ]
    for b_t, b_d in bullets_p1:
        pb1 = tf_p1.add_paragraph()
        pb1.text = f"{b_t} {b_d}"
        pb1.font.name = FONT_BODY
        pb1.font.size = Pt(11.5)
        pb1.font.color.rgb = C_SLATE

    speed_img = os.path.join(ASSETS_DIR, "nano_banana_speed_metrics.png")
    if not os.path.exists(speed_img):
        speed_img = os.path.join(IMAGES_DIR, "high_speed_store_dev_mockup.jpg")
    add_card(s4, Inches(6.9), Inches(1.9), Inches(5.6), Inches(5.0), border_color=C_CARD_BORDER)
    safe_add_image(s4, speed_img, Inches(7.0), Inches(2.0), width=Inches(5.4), height=Inches(4.8))

    # =========================================================
    # SLIDE 5: PILLAR 2 - WHATSAPP COMMERCE & CTWA
    # =========================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, prs)
    add_header(s5, "Pillar 2: Conversational Revenue", "WhatsApp Commerce & Click-to-WhatsApp (CTWA) Marketing", 
               "Meeting customers where they spend 80% of their daily screen time.", tag_color=C_EMERALD)

    wa_cards = [
        ("3-Tap Native Checkout", C_EMERALD, [
            "• In-Chat Catalog browsing without leaving WhatsApp.",
            "• Add-to-cart, item variations, and notes inside chat.",
            "• Razorpay UPI intent links (GPay, PhonePe, Paytm).",
            "• Instant digital invoice & delivery timeline receipt."
        ]),
        ("15-Min Abandoned Cart Recovery", C_GOLD, [
            "• Automated trigger 15 mins after cart abandonment.",
            "• Personalized message with product photo & 5% coupon.",
            "• Recovers up to 68% of lost checkout drop-offs.",
            "• Eliminates friction by continuing existing chat."
        ]),
        ("Click-to-WhatsApp Ads (CTWA)", C_CYAN, [
            "• Meta ads route buyers directly into WhatsApp dialog.",
            "• Instant 15-second response ensures zero ad waste.",
            "• Automated UTM and ad campaign ROI attribution.",
            "• 98% message open rate compared to 12% email."
        ])
    ]
    for k, (w_t, w_col, w_bullets) in enumerate(wa_cards):
        wx = Inches(0.8 + k * 3.95)
        add_card(s5, wx, Inches(1.9), Inches(3.8), Inches(2.5), border_color=w_col)
        tb_w = s5.shapes.add_textbox(wx + Inches(0.25), Inches(2.05), Inches(3.3), Inches(2.2))
        tf_w = tb_w.text_frame
        tf_w.word_wrap = True
        pw_t = tf_w.paragraphs[0]
        pw_t.text = w_t
        pw_t.font.name = FONT_HEADING
        pw_t.font.size = Pt(14)
        pw_t.font.bold = True
        pw_t.font.color.rgb = w_col
        for wb in w_bullets:
            pwb = tf_w.add_paragraph()
            pwb.text = wb
            pwb.font.name = FONT_BODY
            pwb.font.size = Pt(10.5)
            pwb.font.color.rgb = C_SLATE

    img_wa_checkout = os.path.join(IMAGES_DIR, "whatsapp_commerce_checkout_mockup.jpg")
    img_ctwa_dash = os.path.join(IMAGES_DIR, "ctwa_automation_dashboard.jpg")
    
    add_card(s5, Inches(0.8), Inches(4.6), Inches(5.7), Inches(2.3), border_color=C_CARD_BORDER)
    safe_add_image(s5, img_wa_checkout, Inches(0.9), Inches(4.7), width=Inches(5.5), height=Inches(2.1))

    add_card(s5, Inches(6.8), Inches(4.6), Inches(5.7), Inches(2.3), border_color=C_CARD_BORDER)
    safe_add_image(s5, img_ctwa_dash, Inches(6.9), Inches(4.7), width=Inches(5.5), height=Inches(2.1))

    # =========================================================
    # SLIDE 6: PILLAR 3 - AI VOICE AGENTS & COD VERIFICATION
    # =========================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, prs)
    add_header(s6, "Pillar 3: Autonomous Telephony", "Sub-600ms AI Voice Calling: Slashing RTO and Labor Overhead", 
               "Conversational voice agents verifying orders and answering customer inquiries 24/7.", tag_color=C_GOLD)

    add_card(s6, Inches(0.8), Inches(1.9), Inches(5.8), Inches(5.0), border_color=C_GOLD)
    tb_v = s6.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.2), Inches(4.5))
    tf_v = tb_v.text_frame
    tf_v.word_wrap = True

    pv_h = tf_v.paragraphs[0]
    pv_h.text = "THE ZERO-PAYROLL VOICE AGENT"
    pv_h.font.name = FONT_HEADING
    pv_h.font.size = Pt(16)
    pv_h.font.bold = True
    pv_h.font.color.rgb = C_GOLD

    voice_points = [
        ("Instant COD Verification Call:", "Within 60 seconds of a Cash-on-Delivery order, the voice agent calls the buyer, speaks in fluent English or Hindi, and confirms delivery intent."),
        ("Address Sanitization:", "Identifies fake numbers, missing apartment numbers, or mismatched pincodes before dispatching riders."),
        ("RTO Slashing:", "Reduces Return-to-Origin rates from the Indian industry average of 35% down to under 8%."),
        ("Ultra-Low Latency (<600ms):", "No robotic pauses. Realistic conversational interruptions, polite pauses, and human inflection."),
        ("Zero Call Center Salaries:", "Replaces manual calling staff costing ₹25,000/month with zero recurring employee liability.")
    ]
    for v_title, v_body in voice_points:
        pv_b = tf_v.add_paragraph()
        pv_b.text = f"{v_title} {v_body}"
        pv_b.font.name = FONT_BODY
        pv_b.font.size = Pt(11)
        pv_b.font.color.rgb = C_SLATE

    img_voice_dash = os.path.join(IMAGES_DIR, "ai_voice_agent_dashboard.jpg")
    if not os.path.exists(img_voice_dash):
        img_voice_dash = os.path.join(IMAGES_DIR, "whatsapp_calling_dashboard.jpg")
    add_card(s6, Inches(6.9), Inches(1.9), Inches(5.6), Inches(5.0), border_color=C_CARD_BORDER)
    safe_add_image(s6, img_voice_dash, Inches(7.0), Inches(2.0), width=Inches(5.4), height=Inches(4.8))

    # =========================================================
    # SLIDE 7: PILLAR 4 - 1-CLICK FLEET DISPATCH
    # =========================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, prs)
    add_header(s7, "Pillar 4: Physical Fulfillment", "1-Click Hyperlocal Fleet Dispatch & Zero-Payroll Couriers", 
               "Direct API integration with Borzo Business and Porter Logistics for instant counter pickup.", tag_color=C_CYAN)

    add_card(s7, Inches(0.8), Inches(1.9), Inches(5.8), Inches(5.0), border_color=C_CYAN)
    tb_f = s7.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.2), Inches(4.5))
    tf_f = tb_f.text_frame
    tf_f.word_wrap = True

    pf_h = tf_f.paragraphs[0]
    pf_h.text = "THE 1-CLICK COUNTER WORKFLOW"
    pf_h.font.name = FONT_HEADING
    pf_h.font.size = Pt(16)
    pf_h.font.bold = True
    pf_h.font.color.rgb = C_CYAN

    fleet_bullets = [
        ("Step 1: Order Reception:", "Merchant gets notified on WhatsApp & Merchant Dashboard with order items."),
        ("Step 2: 5-Minute Packing Buffer:", "Staff clicks 'Dispatch Now'. The system schedules bike arrival in 5-10 minutes, giving the kitchen or staff time to pack safely."),
        ("Step 3: Automated Rider Assignment:", "Borzo Business API (v1.8) assigns the closest rider in the Mumbai radius automatically."),
        ("Step 4: Live Telemetry Tracking:", "Customer and merchant both receive live GPS tracking links in WhatsApp."),
        ("Step 5: Zero Fixed Fleet Overhead:", "No bikes to maintain, no fuel bills, no driver retention headaches. Pay only per completed delivery.")
    ]
    for fb_t, fb_d in fleet_bullets:
        pf_b = tf_f.add_paragraph()
        pf_b.text = f"{fb_t} {fb_d}"
        pf_b.font.name = FONT_BODY
        pf_b.font.size = Pt(11)
        pf_b.font.color.rgb = C_SLATE

    img_dispatch = os.path.join(IMAGES_DIR, "hyperlocal_dispatch_map_3d.jpg")
    if not os.path.exists(img_dispatch):
        img_dispatch = os.path.join(ASSETS_DIR, "hyperlocal_dispatch_map.jpg")
    add_card(s7, Inches(6.9), Inches(1.9), Inches(5.6), Inches(5.0), border_color=C_CARD_BORDER)
    safe_add_image(s7, img_dispatch, Inches(7.0), Inches(2.0), width=Inches(5.4), height=Inches(4.8))

    # =========================================================
    # SLIDE 8: PILLAR 5 - INTELLIGENT OCR & BACK-OFFICE SYNC
    # =========================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, prs)
    add_header(s8, "Pillar 5: Back-Office Automation", "Gemini Vision OCR & Self-Hosted n8n Workflow Automation", 
               "Eliminating 15+ hours of manual data entry every week for accountants and store owners.", tag_color=C_VIOLET)

    ocr_cards = [
        ("AI Invoice Scanner", C_VIOLET, [
            "• Photo or PDF uploaded via WhatsApp or email.",
            "• Gemini Vision OCR extracts Vendor Name, GSTIN, Line Items, CGST/SGST, and Total Amount.",
            "• 99.4% accuracy on handwritten & thermal bills."
        ]),
        ("Instant Tally & Sheets Export", C_GOLD, [
            "• Formats data directly into Tally XML format.",
            "• Real-time append to Google Sheets accounting.",
            "• Automates monthly GST reconciliation (GSTR-2B).",
            "• Zero manual bookkeeping typing errors."
        ]),
        ("Self-Hosted Private n8n Engine", C_CYAN, [
            "• Runs on dedicated secure private cloud.",
            "• Zero per-task execution fees (unlike Zapier).",
            "• End-to-end data privacy: merchant financial data never leaves private infrastructure.",
            "• Connects ERP, inventory, and messaging."
        ])
    ]
    for m, (oc_t, oc_col, oc_bullets) in enumerate(ocr_cards):
        ox = Inches(0.8 + m * 3.95)
        add_card(s8, ox, Inches(1.9), Inches(3.8), Inches(2.5), border_color=oc_col)
        tb_oc = s8.shapes.add_textbox(ox + Inches(0.25), Inches(2.05), Inches(3.3), Inches(2.2))
        tf_oc = tb_oc.text_frame
        tf_oc.word_wrap = True
        poc_t = tf_oc.paragraphs[0]
        poc_t.text = oc_t
        poc_t.font.name = FONT_HEADING
        poc_t.font.size = Pt(14)
        poc_t.font.bold = True
        poc_t.font.color.rgb = oc_col
        for ob in oc_bullets:
            pob = tf_oc.add_paragraph()
            pob.text = ob
            pob.font.name = FONT_BODY
            pob.font.size = Pt(10.5)
            pob.font.color.rgb = C_SLATE

    img_ocr = os.path.join(IMAGES_DIR, "gemini_ocr_engine.png")
    add_card(s8, Inches(0.8), Inches(4.6), Inches(11.733), Inches(2.3), border_color=C_CARD_BORDER)
    safe_add_image(s8, img_ocr, Inches(1.0), Inches(4.7), width=Inches(11.333), height=Inches(2.1))

    # =========================================================
    # SLIDE 9: FULL SYSTEM ARCHITECTURE & FLOWCHART
    # =========================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, prs)
    add_header(s9, "Technical Blueprint", "End-to-End Autonomous Commerce & Fleet Dispatch Data Pipeline", 
               "Real-time event-driven infrastructure connecting buyer discovery to accounting ledger in sub-seconds.", tag_color=C_GOLD)

    flowchart_img = os.path.join(ASSETS_DIR, "pipeline_data_flowchart_nano_banana.png")
    add_card(s9, Inches(0.8), Inches(1.9), Inches(11.733), Inches(5.1), border_color=C_GOLD)
    safe_add_image(s9, flowchart_img, Inches(0.9), Inches(2.0), width=Inches(11.533), height=Inches(4.9))

    # =========================================================
    # SLIDE 10: COMMERCIAL PACKAGES & PRICING
    # =========================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, prs)
    add_header(s10, "Transparent Pricing & ROI", "Turnkey Client Packages & Monetization Structure", 
               "Clear, predictable agency packages tailored for rapid ROI and long-term partnership.", tag_color=C_EMERALD)

    packages = [
        ("STARTER STOREFRONT", "₹14,999", "₹1,499 / mo", C_CYAN, "Cafes, Bakeries, Boutiques", [
            "• Digital catalog up to 30 items",
            "• Instant UPI QR & GPay integration",
            "• Basic WhatsApp order routing",
            "• Manual dispatch assistance",
            "• Standard 48-hr support SLA"
        ]),
        ("HYPERLOCAL HUB", "₹29,999", "₹3,499 / mo", C_GOLD, "Retail, D2C Brands, Pharmacies (Most Popular)", [
            "• Full headless e-commerce storefront",
            "• Native 3-tap WhatsApp Catalog & Cart",
            "• 1-Click Borzo & Porter API dispatch",
            "• 15-min abandoned cart recovery",
            "• COD verification voice agent",
            "• Priority WhatsApp concierge support"
        ]),
        ("AUTONOMOUS ENTERPRISE", "₹59,999", "₹9,999 / mo", C_MAGENTA, "Multi-Branch Chains & Wholesalers", [
            "• Multi-location dispatch routing",
            "• Gemini Vision OCR for invoice sync",
            "• Automated Tally XML & ERP sync",
            "• Custom AI voice calling scenarios",
            "• Dedicated account engineering",
            "• Weekly conversion audit & ads scaling"
        ])
    ]

    for n, (pkg_name, setup_fee, retainer_fee, pkg_col, target_aud, pkg_feats) in enumerate(packages):
        p_x = Inches(0.8 + n * 3.95)
        add_card(s10, p_x, Inches(1.9), Inches(3.8), Inches(5.1), border_color=pkg_col)

        tb_pk = s10.shapes.add_textbox(p_x + Inches(0.25), Inches(2.1), Inches(3.3), Inches(1.2))
        tf_pk = tb_pk.text_frame
        tf_pk.word_wrap = True
        pp1 = tf_pk.paragraphs[0]
        pp1.text = pkg_name
        pp1.font.name = FONT_HEADING
        pp1.font.size = Pt(13)
        pp1.font.bold = True
        pp1.font.color.rgb = pkg_col

        pp2 = tf_pk.add_paragraph()
        pp2.text = f"{setup_fee} Setup"
        pp2.font.name = FONT_HEADING
        pp2.font.size = Pt(22)
        pp2.font.bold = True
        pp2.font.color.rgb = C_WHITE

        pp3 = tf_pk.add_paragraph()
        pp3.text = f"+ {retainer_fee}"
        pp3.font.name = FONT_HEADING
        pp3.font.size = Pt(12)
        pp3.font.bold = True
        pp3.font.color.rgb = C_GOLD

        add_card(s10, p_x + Inches(0.25), Inches(3.4), Inches(3.3), Inches(0.45), border_color=pkg_col, fill_color=RGBColor(20, 26, 45))
        tb_tg = s10.shapes.add_textbox(p_x + Inches(0.3), Inches(3.45), Inches(3.2), Inches(0.35))
        tf_tg = tb_tg.text_frame
        ptg = tf_tg.paragraphs[0]
        ptg.text = f"Best For: {target_aud}"
        ptg.font.name = FONT_BODY
        ptg.font.size = Pt(9.5)
        ptg.font.color.rgb = C_WHITE

        tb_pf = s10.shapes.add_textbox(p_x + Inches(0.25), Inches(4.0), Inches(3.3), Inches(2.8))
        tf_pf = tb_pf.text_frame
        tf_pf.word_wrap = True
        for pf_item in pkg_feats:
            pfi = tf_pf.add_paragraph() if tf_pf.paragraphs[0].text else tf_pf.paragraphs[0]
            pfi.text = pf_item
            pfi.font.name = FONT_BODY
            pfi.font.size = Pt(10.5)
            pfi.font.color.rgb = C_SLATE

    # =========================================================
    # SLIDE 11: AGENCY LIFECYCLE & CLIENT ACQUISITION FUNNEL
    # =========================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, prs)
    add_header(s11, "Agency Operations", "Client Acquisition Funnel & The 'Zero-Resistance' Offer", 
               "How Bhagat and Yash generate leads, pitch risk-free trials, and lock in recurring retainers.", tag_color=C_CYAN)

    add_card(s11, Inches(0.8), Inches(1.9), Inches(5.8), Inches(5.0), border_color=C_CYAN)
    tb_fn = s11.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.2), Inches(4.5))
    tf_fn = tb_fn.text_frame
    tf_fn.word_wrap = True

    pfn_h = tf_fn.paragraphs[0]
    pfn_h.text = "THE ZERO-RESISTANCE CLIENT CLOSING FORMULA"
    pfn_h.font.name = FONT_HEADING
    pfn_h.font.size = Pt(14)
    pfn_h.font.bold = True
    pfn_h.font.color.rgb = C_CYAN

    fn_steps = [
        ("Phase 1: Local Lead Sourcing", "Target high-intent retail, cafes, pharmacies, and real estate in Mumbai (Bandra, Opera House, Thane, Andheri)."),
        ("Phase 2: The 'Zero-Resistance' Pitch", "'Most stores lose 40% ad spend on slow responses. We will install our instant WhatsApp & Dispatch automation for 3 days on a test run for free. If it works, we partner monthly. If not, you owe ₹0.'"),
        ("Phase 3: 24–48 Hour Deployment", "Zero delay. We plug in verified catalog, WhatsApp Cloud API, and Borzo rider webhook within 48 hours."),
        ("Phase 4: Retainer Conversion", "After seeing real recovered sales and zero dispatch headaches in 3 days, merchants eagerly sign monthly retainers.")
    ]
    for fn_t, fn_d in fn_steps:
        pfn_b = tf_fn.add_paragraph()
        pfn_b.text = f"• {fn_t}: {fn_d}"
        pfn_b.font.name = FONT_BODY
        pfn_b.font.size = Pt(11)
        pfn_b.font.color.rgb = C_SLATE

    funnel_img = os.path.join(ASSETS_DIR, "flowchart_client_funnel.png")
    if not os.path.exists(funnel_img):
        funnel_img = os.path.join(ASSETS_DIR, "flowchart_agency_lifecycle.png")
    add_card(s11, Inches(6.9), Inches(1.9), Inches(5.6), Inches(5.0), border_color=C_CARD_BORDER)
    safe_add_image(s11, funnel_img, Inches(7.0), Inches(2.0), width=Inches(5.4), height=Inches(4.8))

    # =========================================================
    # SLIDE 12: PROVEN ROI & BUSINESS IMPACT METRICS
    # =========================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12, prs)
    add_header(s12, "Concrete ROI", "Proven Performance Benchmarks & Business Impact", 
               "Measurable revenue growth, labor cost reduction, and logistics efficiency.", tag_color=C_EMERALD)

    stat_cards = [
        ("3.4x", "CONVERSION RATE LIFT", "Compared to standard mobile web storefronts", C_CYAN),
        ("98%", "WHATSAPP OPEN RATE", "Within 15 minutes of message dispatch", C_EMERALD),
        ("82%", "DISPATCH TIME SAVED", "From 30+ mins manual booking to 1 click", C_GOLD),
        ("₹0", "FIXED DRIVER SALARIES", "100% on-demand Borzo & Porter logistics", C_MAGENTA)
    ]
    for q, (st_num, st_label, st_desc, st_color) in enumerate(stat_cards):
        qx = Inches(0.8 + q * 2.95)
        add_card(s12, qx, Inches(1.9), Inches(2.8), Inches(1.6), border_color=st_color)
        tb_st = s12.shapes.add_textbox(qx + Inches(0.15), Inches(2.0), Inches(2.5), Inches(1.4))
        tf_st = tb_st.text_frame
        tf_st.word_wrap = True
        pst1 = tf_st.paragraphs[0]
        pst1.text = st_num
        pst1.font.name = FONT_HEADING
        pst1.font.size = Pt(28)
        pst1.font.bold = True
        pst1.font.color.rgb = st_color
        pst2 = tf_st.add_paragraph()
        pst2.text = st_label
        pst2.font.name = FONT_HEADING
        pst2.font.size = Pt(10)
        pst2.font.bold = True
        pst2.font.color.rgb = C_WHITE
        pst3 = tf_st.add_paragraph()
        pst3.text = st_desc
        pst3.font.name = FONT_BODY
        pst3.font.size = Pt(9)
        pst3.font.color.rgb = C_SLATE

    hero_dash = os.path.join(IMAGES_DIR, "hero_revenue_profit_dashboard.jpg")
    if not os.path.exists(hero_dash):
        hero_dash = os.path.join(ASSETS_DIR, "hero_revenue_profit_dashboard.jpg")
    add_card(s12, Inches(0.8), Inches(3.7), Inches(11.733), Inches(3.2), border_color=C_CARD_BORDER)
    safe_add_image(s12, hero_dash, Inches(0.9), Inches(3.8), width=Inches(11.533), height=Inches(3.0))

    # =========================================================
    # SLIDE 13: LEADERSHIP, TECH PARTNERS & 3-DAY PILOT CTA
    # =========================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13, prs)
    add_header(s13, "Leadership & Getting Started", "Deploy Your 3-Day Risk-Free Autonomous Commerce Test Run", 
               "Partner with Mumbai's premier automation and AI commerce engineering lab.", tag_color=C_GOLD)

    # Left: Founders Card & Photo
    add_card(s13, Inches(0.8), Inches(1.9), Inches(5.8), Inches(5.0), border_color=C_GOLD)
    founders_img = os.path.join(IMAGES_DIR, "founders_bhagat_yash_tech_duo.jpg")
    safe_add_image(s13, founders_img, Inches(1.0), Inches(2.1), width=Inches(5.4), height=Inches(2.4))

    tb_ld = s13.shapes.add_textbox(Inches(1.0), Inches(4.6), Inches(5.4), Inches(2.1))
    tf_ld = tb_ld.text_frame
    tf_ld.word_wrap = True
    pld1 = tf_ld.paragraphs[0]
    pld1.text = "Co-Founders: Bhagat Mukhiya & Yash Gaikwad"
    pld1.font.name = FONT_HEADING
    pld1.font.size = Pt(14)
    pld1.font.bold = True
    pld1.font.color.rgb = C_GOLD

    pld2 = tf_ld.add_paragraph()
    pld2.text = "Flying Cheetah Labs was founded to democratize enterprise-grade AI automation for Indian businesses. We engineer high-speed commerce, sub-600ms voice agents, and autonomous logistics dispatch that drive tangible bottom-line profits."
    pld2.font.name = FONT_BODY
    pld2.font.size = Pt(11)
    pld2.font.color.rgb = C_SLATE

    # Right: Call-To-Action Card
    add_card(s13, Inches(6.9), Inches(1.9), Inches(5.6), Inches(5.0), border_color=C_EMERALD)
    tb_cta = s13.shapes.add_textbox(Inches(7.2), Inches(2.1), Inches(5.0), Inches(2.2))
    tf_cta = tb_cta.text_frame
    tf_cta.word_wrap = True

    pcta1 = tf_cta.paragraphs[0]
    pcta1.text = "START YOUR 3-DAY FREE PILOT"
    pcta1.font.name = FONT_HEADING
    pcta1.font.size = Pt(17)
    pcta1.font.bold = True
    pcta1.font.color.rgb = C_EMERALD

    cta_items = [
        "1. Discovery Audit: Analyze ad spend and courier bottlenecks.",
        "2. 48-Hour Staging: Deploy digital catalog & WhatsApp commerce.",
        "3. 3-Day Test Run: Live automated orders & zero dispatch delay.",
        "4. Risk-Free Guarantee: If it doesn't boost profit, you pay ₹0."
    ]
    for ci in cta_items:
        pcta_i = tf_cta.add_paragraph()
        pcta_i.text = ci
        pcta_i.font.name = FONT_BODY
        pcta_i.font.size = Pt(10.5)
        pcta_i.font.color.rgb = C_WHITE

    # Partner Logos Dock (Borzo, Meta, Porter, Razorpay, Shopify)
    add_card(s13, Inches(7.1), Inches(4.25), Inches(5.2), Inches(0.9), border_color=C_CARD_BORDER, fill_color=RGBColor(12, 16, 28))
    partner_logos = ['logo_meta_bsp.png', 'logo_shopify_partner.png', 'logo_razorpay.png', 'logo_borzo.png', 'logo_porter.png']
    for idx, pl_name in enumerate(partner_logos):
        pl_path = os.path.join(IMAGES_DIR, pl_name)
        safe_add_image(s13, pl_path, Inches(7.2 + idx * 1.0), Inches(4.35), width=Inches(0.9), height=Inches(0.68))

    # Contact Info Pill
    add_card(s13, Inches(7.1), Inches(5.3), Inches(5.2), Inches(1.3), border_color=C_GOLD, fill_color=RGBColor(20, 26, 42))
    tb_con = s13.shapes.add_textbox(Inches(7.25), Inches(5.38), Inches(4.9), Inches(1.1))
    tf_con = tb_con.text_frame
    tf_con.word_wrap = True
    pcon1 = tf_con.paragraphs[0]
    pcon1.text = "READY TO AUTOMATE YOUR STORE?"
    pcon1.font.name = FONT_HEADING
    pcon1.font.size = Pt(11)
    pcon1.font.bold = True
    pcon1.font.color.rgb = C_GOLD

    pcon2 = tf_con.add_paragraph()
    pcon2.text = "📞 Phone / WhatsApp: +91 99203 86421\n🌐 Live Web Portal: https://flying-cheetah.vercel.app/\n📍 Mumbai: Opera House / Bandra West"
    pcon2.font.name = FONT_BODY
    pcon2.font.size = Pt(10.5)
    pcon2.font.color.rgb = C_WHITE

    # Save Presentation
    prs.save(OUTPUT_PPTX)
    print(f"Master presentation successfully generated at: {OUTPUT_PPTX}")

if __name__ == "__main__":
    build_presentation()
